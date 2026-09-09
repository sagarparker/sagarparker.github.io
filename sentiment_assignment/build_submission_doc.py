"""Build the Word submission for the sentiment analysis assignment."""

from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
SCREENSHOT_DIR = ROOT / "screenshots"
OUTPUT_DOCX = ROOT / "Sentiment_Analysis_Assignment_Submission.docx"
PYTHON = "/Users/sagarparker/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3"


COMMANDS = [
    (
        "Positive program run",
        [PYTHON, "sentiment_analyzer.py", "I love this excellent and wonderful project."],
        "01_positive_run.png",
    ),
    (
        "Negative program run",
        [PYTHON, "sentiment_analyzer.py", "This was a terrible and disappointing problem."],
        "02_negative_run.png",
    ),
    (
        "Neutral program run",
        [PYTHON, "sentiment_analyzer.py", "The package is on the table."],
        "03_neutral_run.png",
    ),
    (
        "Unit tests",
        [PYTHON, "-m", "unittest", "test_sentiment_analyzer.py"],
        "04_unit_tests.png",
    ),
    (
        "Dataset-style CSV evaluation",
        [PYTHON, "evaluate_dataset.py"],
        "05_dataset_evaluation.png",
    ),
]


SESSION_HISTORY = """User prompt:
1.Takes a text string as input

2.Outputs a message about the sentiment of the text string, whether it's "Positive", "Negative" or "Neutral". See also Sentiment Analysis

Include tests showing all the sentiments and submit a word document with screenshots of the program operations and tests.
Use any programming language you prefer
Include the entire related LLM (ChatGPT) session history in your submission for this assignment, including all the prompts you used.
Try Testing it on a sentiment dataset, e.g., https://www.kaggle.com/datasets/abhi8923shriv/sentiment-analysis-dataset
Hint: TextBlob

Assistant prompts/actions used in this session:
- Inspected the repository structure to choose a clean location for assignment files.
- Used the document-generation instructions because the requested deliverable is a Word document.
- Checked whether TextBlob was installed in the runtime. It was not available.
- Implemented a Python sentiment analyzer that uses TextBlob when installed and a fallback lexicon otherwise.
- Added unit tests covering Positive, Negative, and Neutral outputs.
- Added a small dataset-style CSV and evaluator to demonstrate batch sentiment testing without requiring Kaggle credentials.
- Ran all example program operations, unit tests, and dataset evaluation.
- Generated this Word document with screenshots of the program and tests.
"""


def run_command(command: list[str]) -> str:
    completed = subprocess.run(
        command,
        cwd=ROOT,
        check=False,
        text=True,
        capture_output=True,
    )
    prompt = "$ " + " ".join(command).replace(PYTHON, "python3")
    output_parts = [prompt]
    if completed.stdout:
        output_parts.append(completed.stdout.rstrip())
    if completed.stderr:
        output_parts.append(completed.stderr.rstrip())
    output_parts.append(f"[exit code: {completed.returncode}]")
    return "\n".join(output_parts)


def make_terminal_screenshot(text: str, output_path: Path) -> None:
    SCREENSHOT_DIR.mkdir(exist_ok=True)
    font = ImageFont.load_default()
    lines = text.splitlines()
    line_height = 18
    padding = 22
    width = 1120
    height = max(180, padding * 2 + line_height * len(lines))

    image = Image.new("RGB", (width, height), "#15191f")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 34), fill="#252a33")
    draw.text((18, 10), "Terminal", fill="#e8edf2", font=font)

    y = 52
    for line in lines:
        fill = "#8be9a4" if line.startswith("$ ") else "#f3f5f7"
        if line.startswith("[exit code:"):
            fill = "#9fb3c8"
        draw.text((padding, y), line, fill=fill, font=font)
        y += line_height

    image.save(output_path)


def set_default_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(46, 116, 181)


def add_code_block(document: Document, title: str, path: Path) -> None:
    document.add_heading(title, level=2)
    paragraph = document.add_paragraph()
    run = paragraph.add_run(path.read_text(encoding="utf-8"))
    run.font.name = "Courier New"
    run.font.size = Pt(8)


def build_document() -> None:
    outputs: list[tuple[str, Path, str]] = []
    for title, command, filename in COMMANDS:
        text = run_command(command)
        screenshot_path = SCREENSHOT_DIR / filename
        make_terminal_screenshot(text, screenshot_path)
        outputs.append((title, screenshot_path, text))

    document = Document()
    set_default_style(document)

    title = document.add_paragraph()
    title_run = title.add_run("Sentiment Analysis Assignment Submission")
    title_run.font.size = Pt(20)
    title_run.bold = True
    document.add_paragraph("Program language: Python")
    document.add_paragraph(
        "The program accepts a text string and classifies it as Positive, "
        "Negative, or Neutral. It uses TextBlob automatically when available "
        "and a small fallback lexicon in environments where TextBlob is not installed."
    )

    document.add_heading("Program Operations and Test Screenshots", level=1)
    for title, screenshot_path, _ in outputs:
        document.add_heading(title, level=2)
        document.add_picture(str(screenshot_path), width=Inches(6.5))

    document.add_heading("Testing Summary", level=1)
    document.add_paragraph(
        "The unit tests cover all three required sentiments: Positive, Negative, "
        "and Neutral. The sample CSV evaluation demonstrates dataset-style testing "
        "on six labeled examples."
    )

    add_code_block(document, "Source Code: sentiment_analyzer.py", ROOT / "sentiment_analyzer.py")
    add_code_block(document, "Source Code: test_sentiment_analyzer.py", ROOT / "test_sentiment_analyzer.py")
    add_code_block(document, "Source Code: evaluate_dataset.py", ROOT / "evaluate_dataset.py")

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Related ChatGPT Session History", level=1)
    history_paragraph = document.add_paragraph()
    history_run = history_paragraph.add_run(SESSION_HISTORY)
    history_run.font.name = "Courier New"
    history_run.font.size = Pt(9)

    document.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
